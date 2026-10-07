#!/usr/bin/env python3
"""AppFolio export loader + feature builder for PropertyManagement.com member skills.

Reads the standard AppFolio report exports (xlsx or csv) from a folder, detects each report
by its column signature, normalizes the rows, and writes:

  out/inventory.json        which reports were found, row counts, as-of dates, what is missing
  out/company.json          company-wide metrics (occupancy, vacancy, departures, renewals,
                            collections, turns, delinquency, by region / manager / property type)
  out/owners.json           one record per property owner: portfolio, signals, risk score + tier
  out/data-quality.md       plain-English notes on gaps that change how numbers should be read

Usage
  python3 appfolio_loader.py inventory <exports_dir>
  python3 appfolio_loader.py features  <exports_dir> --out <out_dir> [--survey survey.json]
                                       [--as-of YYYY-MM-DD] [--config client-config.json]

Only the standard library plus openpyxl are required. Every report is optional; a missing
report leaves its features null and is listed in data-quality.md. Never invents a number.
"""
import argparse, collections, csv, datetime as dt, json, os, re, statistics, sys
try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None

# ----------------------------------------------------------------------------- helpers
def isd(v): return isinstance(v, (dt.datetime, dt.date))
def to_dt(v):
    if isinstance(v, dt.datetime): return v
    if isinstance(v, dt.date): return dt.datetime(v.year, v.month, v.day)
    if isinstance(v, str):
        s = v.strip()
        for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%Y-%m-%d %H:%M:%S", "%m/%d/%y"):
            try: return dt.datetime.strptime(s[:19] if "%H" in fmt else s[:10], fmt)
            except ValueError: pass
    return None
def num(v):
    if v in (None, ""): return None
    if isinstance(v, (int, float)): return float(v)
    s = str(v).replace("$", "").replace(",", "").strip()
    if s.startswith("(") and s.endswith(")"): s = "-" + s[1:-1]
    try: return float(s)
    except ValueError: return None
def sid(v):
    """Vendor ids arrive as int, float or str; normalize to a string key."""
    if v in (None, ""): return None
    if isinstance(v, float) and v.is_integer(): return str(int(v))
    if isinstance(v, int): return str(v)
    s = str(v).strip()
    return s[:-2] if re.fullmatch(r"\d+\.0", s) else s
def med(xs):
    xs = [x for x in xs if x is not None]
    return round(statistics.median(xs), 1) if xs else None
def pct(n, d, nd=1): return round(100.0 * n / d, nd) if d else None
TEST_RE = re.compile(r"\btest\b|zdummy|dummy", re.I)

ZW_RE = re.compile(r"[\u200b\u200c\u200d\ufeff]")   # AppFolio directories carry invisible zero-width characters on some names
def clean(v): return re.sub(r"\s+", " ", ZW_RE.sub("", str(v or ""))).strip()
ENTITY_RE = re.compile(r"\b(LLC|L\.L\.C\.?|Inc\.?|Corp\.?|Co\.|Company|Trust|Properties|Property|Holdings|Investments?|Ventures|Capital|Partners|Group|Fund|Estates?|Realty|Management|Development|Enterprises|Rentals|Homes)\b", re.I)
def pretty_name(v):
    """'Breslin, Daniel' -> 'Daniel Breslin'; entity names and anything with more than one comma are left alone, and trust boilerplate is cut at the first clause."""
    v = clean(v)
    if ENTITY_RE.search(v) or v.count(",") != 1: return re.split(r",? as trustee| under trust| dated ", v, maxsplit=1)[0].strip(" ,")
    last, first = [x.strip() for x in v.split(",")]
    return f"{first} {last}".strip()

def first_name_of(v):
    """'Viet (Scott)' -> 'Scott': a nickname in parentheses is what the person is called."""
    v = clean(v); m = re.search(r"\(([^)]+)\)", v)
    return (m.group(1) if m else v).strip().title()

OWNER_EXIT_RE = re.compile(r"management company terminated|owner selling|owner moved|property sold|owner cancel|owner terminat|selling as is|cannot afford", re.I)

def read_table(path):
    """Return (header_row_index, headers, rows) for an AppFolio export. Header = first row with >= 5 filled cells."""
    rows = []
    if path.lower().endswith(".csv"):
        with open(path, newline="", encoding="utf-8-sig", errors="replace") as f:
            rows = [r for r in csv.reader(f)]
    else:
        if openpyxl is None: raise SystemExit("openpyxl is required for .xlsx exports (pip install openpyxl)")
        ws = openpyxl.load_workbook(path, read_only=True, data_only=True).worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    hi = next((i for i, r in enumerate(rows) if r and sum(1 for v in r if v not in (None, "")) >= 5), None)
    if hi is None: return None, [], []
    hdr = [str(v).strip() if v is not None else "" for v in rows[hi]]
    data = [list(r) + [None] * (len(hdr) - len(r)) for r in rows[hi + 1:] if r and any(v not in (None, "") for v in r)]
    return hi, hdr, data

SIGNATURES = {  # report -> columns that identify it (all must be present)
    "rent_roll": ["Unit", "Status", "Market Rent", "Rent", "Lease To", "Past Due"],
    "tenant_tickler": ["Event", "Date", "Move Out Reason"],
    "lease_history": ["Lease End", "Outcome", "Tenant ID"],
    "renewal_summary": ["Previous Lease End", "Countersigned Date", "Percent Difference"],
    "lease_expiration": ["Lease Expires", "Last Lease Renewal", "Owner(s)"],
    "property_directory": ["Property ID", "Site Manager Name", "Management End Date"],
    "owner_directory": ["Owner ID", "Properties Owned IDs", "Email"],
    "unit_directory": ["Unit Name", "Rent Ready", "Unit ID", "Bedrooms"],
    "work_orders": ["Work Order Number", "Work Done On", "Created At"],
    "charge_detail": ["Charge Amount", "Paid Amount", "Account Number"],
    "delinquency": ["Amount Receivable", "Delinquent Rent", "Tenant Status"],
    "unit_vacancy_detail": ["Days Vacant", "Unit Status", "Rent Ready"],
}
def detect(hdr):
    hs = set(hdr)
    for name, cols in SIGNATURES.items():
        if all(c in hs for c in cols): return name
    return None

def exported_on(path):
    """AppFolio prints 'Exported On: MM/DD/YYYY ...' in the preamble; fall back to file mtime."""
    try:
        if path.lower().endswith(".csv"):
            with open(path, encoding="utf-8-sig", errors="replace") as f: head = f.read(2000)
        else:
            ws = openpyxl.load_workbook(path, read_only=True, data_only=True).worksheets[0]
            head = " ".join(str(c) for r in ws.iter_rows(min_row=1, max_row=8, values_only=True) for c in r if c)
        m = re.search(r"Exported On:\s*(\d{2}/\d{2}/\d{4})", head)
        if m: return dt.datetime.strptime(m.group(1), "%m/%d/%Y").date().isoformat()
    except Exception: pass
    return dt.date.fromtimestamp(os.path.getmtime(path)).isoformat()

class Report:
    def __init__(self, kind, path, hdr, rows):
        self.kind, self.path, self.hdr, self.rows = kind, path, hdr, rows
        self.ix = {h: i for i, h in enumerate(hdr)}
        self.as_of = exported_on(path)
    def col(self, r, name): 
        i = self.ix.get(name); return r[i] if i is not None and i < len(r) else None
    def has(self, name): return name in self.ix

def load_dir(d):
    reports = {}
    for fn in sorted(os.listdir(d)):
        p = os.path.join(d, fn)
        if not fn.lower().endswith((".xlsx", ".xlsm", ".csv")) or fn.startswith("~$"): continue
        try: hi, hdr, rows = read_table(p)
        except Exception as ex:
            print(f"skip {fn}: {ex!r}", file=sys.stderr); continue
        kind = detect(hdr) if hdr else None
        if not kind: print(f"unrecognized export (no signature match): {fn}", file=sys.stderr); continue
        rep = Report(kind, p, hdr, rows)
        if kind in reports:  # keep the newest export of a kind
            if rep.as_of >= reports[kind].as_of: reports[kind] = rep
        else: reports[kind] = rep
    return reports

# ----------------------------------------------------------------------------- features
def region_of(city, cfg):
    c = (city or "").strip().lower()
    if not c: return "Unknown"
    core = [x.lower() for x in cfg.get("regions", {}).get("core_cities", [])]
    return cfg.get("regions", {}).get("core_label", "City") if c in core else cfg.get("regions", {}).get("other_label", "Suburbs")

def build(reports, cfg, survey, as_of):
    TODAY = to_dt(as_of) or dt.datetime.now()
    dq = []  # data-quality notes
    props, owners, out_owners = {}, {}, {}

    # --- property directory: the spine (type, city, manager, active, owner ids)
    pd_ = reports.get("property_directory")
    if pd_:
        for r in pd_.rows:
            p = sid(pd_.col(r, "Property ID"))
            if not p or TEST_RE.search(str(pd_.col(r, "Property") or "")): continue
            end = to_dt(pd_.col(r, "Management End Date")); start = to_dt(pd_.col(r, "Management Start Date"))
            mgr = str(pd_.col(r, "Site Manager Name") or "").strip()
            props[p] = {"id": p, "name": clean(pd_.col(r, "Property Name") or pd_.col(r, "Property")),
                        "type": str(pd_.col(r, "Property Type") or "").strip(), "city": str(pd_.col(r, "Property City") or "").strip(),
                        "address": clean(pd_.col(r, "Property Street Address 1")),   # street line only; owners recognise their property by it
                        "zip": str(pd_.col(r, "Property Zip") or "").strip()[:5], "manager": (re.split(r"[\s,]+", mgr)[0].title() if mgr else "Unassigned"),
                        "units_listed": num(pd_.col(r, "Units")), "active": (end is None) or end > TODAY,
                        "management_start": start.date().isoformat() if start else None, "management_end": end.date().isoformat() if end else None,
                        "owner_ids": [sid(x) for x in str(pd_.col(r, "Owner ID(s)") or "").split(",") if x.strip()]}
        blank_mgr = sum(1 for p in props.values() if p["active"] and p["manager"] == "Unassigned")
        if blank_mgr: dq.append(f"Site Manager Name is blank on {blank_mgr} active properties; those owners cannot be attributed to a manager.")
    else:
        dq.append("No Property Directory export: no manager, property type, or management end dates. Owner attribution runs on the rent roll only.")
    def prop_active(p): return props.get(p, {}).get("active", True)

    # --- owner directory
    od = reports.get("owner_directory")
    if od:
        for r in od.rows:
            oid = sid(od.col(r, "Owner ID"))
            if not oid: continue
            emails = [x.strip().lower() for x in re.split(r"[,;\s]+", str(od.col(r, "Email") or "")) if "@" in x]
            pids = [sid(x) for x in str(od.col(r, "Properties Owned IDs") or "").split(",") if x.strip()]
            owners[oid] = {"owner_id": oid, "name": clean(od.col(r, "Name")), "first_name": first_name_of(od.col(r, "First Name")),
                           "emails": emails, "property_ids": [p for p in pids if p and prop_active(p)], "all_property_ids": [p for p in pids if p],
                           "send_packets_by_email": str(od.col(r, "Send Owner Packets by Email?") or ""),
                           "hold_payments": str(od.col(r, "Hold Payments?") or od.col(r, "Hold Payments") or "").strip().lower() in ("yes", "y", "true"),
                           "last_payment": to_dt(od.col(r, "Last Payment Date")), "tags": clean(od.col(r, "Tags"))}
    else:
        dq.append("No Owner Directory export: survey responses cannot be joined to properties, and owner-level output is unavailable.")

    # --- rent roll: units per property
    rr = reports.get("rent_roll"); units_by_prop = collections.defaultdict(list)
    OCC = ("Current", "Notice-Rented", "Notice-Unrented", "Evict"); UNIT_STATUSES = OCC + ("Vacant-Unrented", "Vacant-Rented")
    if rr:
        for r in rr.rows:
            if rr.col(r, "Unit ID") in (None, "") or str(rr.col(r, "Status") or "") not in UNIT_STATUSES: continue
            if TEST_RE.search(str(rr.col(r, "Property") or "") + str(rr.col(r, "Tenant") or "")): continue
            p = sid(rr.col(r, "Property ID"))
            if not prop_active(p): continue
            units_by_prop[p].append({"unit": clean(rr.col(r, "Unit")), "status": str(rr.col(r, "Status")), "move_in": to_dt(rr.col(r, "Move-in")) if rr.has("Move-in") else None, "rent": num(rr.col(r, "Rent")), "market_rent": num(rr.col(r, "Computed Market Rent")) or num(rr.col(r, "Market Rent")),
                                     "past_due": num(rr.col(r, "Past Due")) or 0.0, "lease_to": to_dt(rr.col(r, "Lease To")), "last_move_out": to_dt(rr.col(r, "Last Move Out")) if rr.has("Last Move Out") else None,
                                     "rent_ready": str(rr.col(r, "Rent Ready") or ""), "type": str(rr.col(r, "Property Type") or props.get(p, {}).get("type", "")), "city": str(rr.col(r, "Property City") or props.get(p, {}).get("city", ""))})
        if not rr.has("Property ID"): dq.append("Rent roll lacks the Property ID column; export with the property columns selected so units join to owners.")
    else:
        dq.append("No Rent Roll export: occupancy, rent vs market and past-due balances are unavailable.")

    # --- tenant tickler: departures
    tt = reports.get("tenant_tickler"); moveouts = []
    if tt:
        for r in tt.rows:
            d = to_dt(tt.col(r, "Date"))
            if not d or str(tt.col(r, "Event")) != "Move-out" or TEST_RE.search(str(tt.col(r, "Property") or "")): continue
            reason = str(tt.col(r, "Move Out Reason") or "").strip()
            moveouts.append({"date": d, "property_id": sid(tt.col(r, "Property ID")), "reason": reason or "No reason recorded", "owner_initiated": bool(OWNER_EXIT_RE.search(reason)),
                             "city": str(tt.col(r, "Property City") or ""), "user": str(tt.col(r, "Assigned User") or "").strip()})
        movein_n = sum(1 for r in tt.rows if str(tt.col(r, "Event")) == "Move-in" and to_dt(tt.col(r, "Date")))
        blank = sum(1 for m in moveouts if m["reason"] == "No reason recorded")
        if moveouts and blank / len(moveouts) > 0.3: dq.append(f"{blank} of {len(moveouts)} move-outs carry no Move Out Reason, so owner-initiated departures are a floor, not a count.")
    else:
        movein_n = None; dq.append("No Tenant Tickler export: departures and owner-initiated exits are unavailable.")

    # --- lease history: renewal outcomes
    lh = reports.get("lease_history"); coh = []
    if lh:
        ends = [to_dt(lh.col(r, "Lease End")) for r in lh.rows]; ends = [e for e in ends if e]
        w0 = min(ends) if ends else None
        for r in lh.rows:
            le = to_dt(lh.col(r, "Lease End")); o = str(lh.col(r, "Outcome") or "")
            if not le or o not in ("Renewed - Fixed Term", "Moved Out", "Month-to-Month") or TEST_RE.search(str(lh.col(r, "Property") or "")): continue
            mo = to_dt(lh.col(r, "Move Out"))
            if o == "Moved Out" and mo and w0 and mo < w0 - dt.timedelta(days=45): continue  # prior-period lease break
            coh.append({"outcome": o, "renewed": o != "Moved Out", "property_id": sid(lh.col(r, "Property ID")), "city": str(lh.col(r, "Property City") or ""), "lease_end": le})
    else:
        dq.append("No Lease History export: renewal outcomes by property and region are unavailable.")

    # --- renewal summary: pricing on renewals
    rs = reports.get("renewal_summary"); ren_pct = []; ren_flat = 0; ren_n = 0
    if rs:
        for r in rs.rows:
            if str(rs.col(r, "Status")) != "Renewed": continue
            ren_n += 1; p = num(rs.col(r, "Percent Difference"))
            if p is not None: ren_pct.append(p); ren_flat += 1 if p <= 0 else 0

    # --- lease expiration detail: upcoming expirations + m2m + owner column
    le_ = reports.get("lease_expiration"); expiring_by_prop = collections.defaultdict(list); m2m_by_prop = collections.Counter(); m2m_units_by_prop = collections.defaultdict(list)
    if le_:
        for r in le_.rows:
            if le_.col(r, "Unit ID") in (None, "") or TEST_RE.search(str(le_.col(r, "Property") or "")): continue
            p = sid(le_.col(r, "Property ID")); exp = to_dt(le_.col(r, "Lease Expires"))
            if exp is None: m2m_by_prop[p] += 1; m2m_units_by_prop[p].append(clean(le_.col(r, "Unit"))); continue
            if TODAY <= exp <= TODAY + dt.timedelta(days=120):
                expiring_by_prop[p].append({"unit": clean(le_.col(r, "Unit")), "status": str(le_.col(r, "Status") or ""), "expires": exp, "rent": num(le_.col(r, "Rent")), "market_rent": num(le_.col(r, "Computed Market Rent")) or num(le_.col(r, "Market Rent")),
                                            "renewal_started": to_dt(le_.col(r, "Renewal Start Date")) is not None, "notice_given": to_dt(le_.col(r, "Notice Given Date")) is not None})
    # --- unit vacancy detail
    uv = reports.get("unit_vacancy_detail"); vac_by_prop = collections.defaultdict(list)
    if uv:
        for r in uv.rows:
            st = str(uv.col(r, "Unit Status") or "")
            if not st.startswith("Vacant") or TEST_RE.search(str(uv.col(r, "Property") or "")): continue
            p = sid(uv.col(r, "Property ID"))
            if not prop_active(p): continue
            vac_by_prop[p].append({"unit": clean(uv.col(r, "Unit") or uv.col(r, "Unit Name")), "status": st, "days_vacant": num(uv.col(r, "Days Vacant")), "rent_ready": str(uv.col(r, "Rent Ready") or "") == "Yes", "posted": str(uv.col(r, "Posted To Website") or "") == "Yes",
                                   "advertised_rent": num(uv.col(r, "Advertised Rent")), "last_move_out": to_dt(uv.col(r, "Last Move Out"))})
        zero_days = sum(1 for v in vac_by_prop.values() for u in v if u["days_vacant"] == 0)
        if zero_days: dq.append(f"Days Vacant reads 0 on {zero_days} vacant units (no move-out or available-on date recorded); vacancy age is understated there.")
    # --- work orders (turns / resident / internal)
    wo = reports.get("work_orders"); turns_by_prop = collections.defaultdict(list); wo_types = collections.Counter(); est_pairs = []
    if wo:
        for r in wo.rows:
            t = str(wo.col(r, "Work Order Type") or ""); c = to_dt(wo.col(r, "Created At"))
            if not t or not c: continue
            wo_types[t] += 1; p = sid(wo.col(r, "Property ID")); done = to_dt(wo.col(r, "Work Done On")); st = str(wo.col(r, "Status") or "")
            note = str(wo.col(r, "Status Notes") or "") if wo.has("Status Notes") else ""
            est, act = num(wo.col(r, "Estimate Amount")), (num(wo.col(r, "Vendor Bill Amount")) or num(wo.col(r, "Amount")))
            if est and act: est_pairs.append((est, act))
            turns_by_prop[p].append({"type": t, "unit": clean(wo.col(r, "Unit") or wo.col(r, "Unit Name")), "status": st, "days": (done - c).days if (done and (done - c).days >= 0) else None, "canceled": st == "Canceled", "open": done is None and st not in ("Completed", "Canceled", "Closed"),
                                     "owner_exit_note": bool(OWNER_EXIT_RE.search(note)), "vendor": str(wo.col(r, "Vendor") or "").strip(), "created": c})
        if wo_types and "Resident" not in wo_types: dq.append("The work-order export holds no Resident-type orders (turns and internal only); resident maintenance likely lives in another system, so repair speed here is turn speed.")
    # --- charge detail (rent only)
    cd = reports.get("charge_detail"); coll_by_prop = collections.defaultdict(lambda: [0.0, 0.0]); rent_gl = str(cfg.get("rent_gl_account", "4100"))
    if cd:
        for r in cd.rows:
            if str(cd.col(r, "Account Number") or "").strip() != rent_gl: continue
            if cd.has("Party Type") and str(cd.col(r, "Party Type") or "") not in ("Occupancy", ""): continue
            ch = num(cd.col(r, "Charge Amount"))
            if ch is None: continue
            p = sid(cd.col(r, "Property ID")); coll_by_prop[p][0] += ch; coll_by_prop[p][1] += num(cd.col(r, "Paid Amount")) or 0.0
    # --- delinquency
    dq_rep = reports.get("delinquency"); delq_by_prop = collections.defaultdict(lambda: {"tenants": 0, "receivable": 0.0, "evict": 0}); legacy_by_prop = collections.Counter()
    if dq_rep:
        legacy = 0.0; legacy_rows = 0
        for r in dq_rep.rows:
            st = str(dq_rep.col(r, "Tenant Status") or ""); mo = to_dt(dq_rep.col(r, "Move Out"))
            gone = st == "Past" or (st == "Evict" and mo is not None and mo <= TODAY)   # an evicted tenant who has moved out is a former tenant: legacy AR, not a live balance
            if gone: amt = num(dq_rep.col(r, "Amount Receivable")) or 0.0; legacy += amt; legacy_rows += 1; legacy_by_prop[sid(dq_rep.col(r, "Property ID"))] += amt; continue
            if st not in ("Current", "Notice", "Evict"): continue
            p = sid(dq_rep.col(r, "Property ID")); d = delq_by_prop[p]; d["tenants"] += 1; d["receivable"] += num(dq_rep.col(r, "Amount Receivable")) or 0.0; d["evict"] += 1 if st == "Evict" else 0
        if legacy_rows: dq.append(f"{legacy_rows} delinquency rows are former tenants (Past, or Evict with a move-out date behind us): ${legacy:,.0f} of legacy receivable, kept out of the per-owner balance signal.")
        mism = sum(1 for p, us in units_by_prop.items() if abs(sum(u["past_due"] for u in us) - (delq_by_prop[p]["receivable"] if p in delq_by_prop else 0.0)) > 500)
        if mism: dq.append(f"Rent-roll Past Due and Delinquency receivable disagree by more than $500 on {mism} properties (different as-of times or netting); owner cards show the rent-roll figure.")
    else: legacy = None

    # --- survey (optional JSON produced from the PropertyManagement.com connector or a dashboard export)
    surv_by_email = collections.defaultdict(list)
    for s in (survey or []):
        em = (s.get("email") or "").strip().lower()
        if em: surv_by_email[em].append(s)

    # ------------------------------------------------------------- company-wide
    all_units = [u for v in units_by_prop.values() for u in v]
    occ_n = sum(1 for u in all_units if u["status"] in OCC)
    never_leased = [u for u in all_units if u["status"] == "Vacant-Unrented" and rr and rr.has("Last Move Out") and u["last_move_out"] is None]
    leasable = len(all_units) - len(never_leased)
    def occ_split(key):
        out = {}
        for k in sorted({u[key] for u in all_units if u[key]}):
            sub = [u for u in all_units if u[key] == k]; out[k] = {"units": len(sub), "occupancy_pct": pct(sum(1 for u in sub if u["status"] in OCC), len(sub))}
        return out
    mo_recent = moveouts
    company = {
        "as_of": {k: v.as_of for k, v in reports.items()}, "reports_found": sorted(reports.keys()),
        "portfolio": {"active_properties": sum(1 for p in props.values() if p["active"]) if props else None, "properties_ever": len(props) or None,
                       "active_units": len(all_units) or None, "occupied": occ_n if all_units else None, "occupancy_pct_raw": pct(occ_n, len(all_units)),
                       "never_leased_vacant_units": len(never_leased) if all_units else None, "occupancy_pct_leasable": pct(occ_n, leasable),
                       "vacant_unrented": sum(1 for u in all_units if u["status"] == "Vacant-Unrented") if all_units else None,
                       "past_due_total": round(sum(u["past_due"] for u in all_units), 2) if all_units else None,
                       "avg_rent": round(statistics.mean([u["rent"] for u in all_units if u["rent"] and u["status"] in OCC]), 0) if all_units else None,
                       "by_type": occ_split("type") if all_units else None, "by_city": occ_split("city") if all_units else None},
        "departures": {"window_start": min((m["date"] for m in mo_recent), default=None).date().isoformat() if mo_recent else None, "window_end": max((m["date"] for m in mo_recent), default=None).date().isoformat() if mo_recent else None,
                        "move_outs": len(mo_recent) if tt else None, "move_ins": movein_n, "owner_initiated": sum(1 for m in mo_recent if m["owner_initiated"]) if tt else None,
                        "reasons": dict(collections.Counter(m["reason"] for m in mo_recent).most_common()) if tt else None,
                        "owner_initiated_by_region": dict(collections.Counter(region_of(m["city"], cfg) for m in mo_recent if m["owner_initiated"])) if tt else None},
        "renewals": {"cohort": len(coh) or None, "renewed": sum(1 for c in coh if c["renewed"]) if coh else None, "rate_pct": pct(sum(1 for c in coh if c["renewed"]), len(coh)),
                      "by_region": {k: {"n": len(v), "rate_pct": pct(sum(1 for c in v if c["renewed"]), len(v))} for k, v in collections.defaultdict(list, {}).items()} if False else
                                   {reg: {"n": len([c for c in coh if region_of(c["city"], cfg) == reg]), "rate_pct": pct(sum(1 for c in coh if region_of(c["city"], cfg) == reg and c["renewed"]), len([c for c in coh if region_of(c["city"], cfg) == reg]))}
                                    for reg in sorted({region_of(c["city"], cfg) for c in coh})} if coh else None,
                      "by_manager": {m: {"n": n, "rate_pct": pct(rn, n)} for m, (n, rn) in sorted({mm: (len([c for c in coh if props.get(c["property_id"], {}).get("manager") == mm]), sum(1 for c in coh if props.get(c["property_id"], {}).get("manager") == mm and c["renewed"])) for mm in {props.get(c["property_id"], {}).get("manager", "Unassigned") for c in coh}}.items()) if n >= 5} if (coh and props) else None,
                      "renewals_priced": ren_n or None, "median_pct_increase": med(ren_pct), "renewed_flat_or_down": ren_flat if rs else None,
                      "expiring_120d": sum(len(v) for v in expiring_by_prop.values()) if le_ else None, "standing_m2m": sum(m2m_by_prop.values()) if le_ else None,
                      "note": "rate = renewed / (renewed + moved out) on the Lease History window, prior-period lease breaks excluded; departures the tenant did not choose are NOT removed here (see metric-definitions.md)"},
        "vacancy": {"vacant_units": sum(len(v) for v in vac_by_prop.values()) if uv else None, "rent_ready": sum(1 for v in vac_by_prop.values() for u in v if u["rent_ready"]) if uv else None,
                     "posted_to_website": sum(1 for v in vac_by_prop.values() for u in v if u["posted"]) if uv else None,
                     "median_days_vacant_tracked": med([u["days_vacant"] for v in vac_by_prop.values() for u in v if u["days_vacant"] and u["days_vacant"] > 0]) if uv else None,
                     "over_60_days": sum(1 for v in vac_by_prop.values() for u in v if u["days_vacant"] and u["days_vacant"] > 60) if uv else None},
        "turns": {"types": dict(wo_types) if wo else None, "measured": sum(1 for v in turns_by_prop.values() for t in v if t["days"] is not None) if wo else None,
                   "median_days": med([t["days"] for v in turns_by_prop.values() for t in v]) if wo else None, "canceled": sum(1 for v in turns_by_prop.values() for t in v if t["canceled"]) if wo else None,
                   "canceled_owner_exit_notes": sum(1 for v in turns_by_prop.values() for t in v if t["owner_exit_note"]) if wo else None,
                   "by_vendor": sorted(({"vendor": k or "No vendor", "n": len(v), "median_days": med([t["days"] for t in v])} for k, v in collections.defaultdict(list, {t["vendor"]: [x for vv in turns_by_prop.values() for x in vv if x["vendor"] == t["vendor"]] for vv in turns_by_prop.values() for t in vv}).items() if v), key=lambda d: -d["n"])[:6] if wo else None,
                   "estimate_vs_actual": {"pairs": len(est_pairs), "median_variance_pct": med([100 * (a - e) / e for e, a in est_pairs if e]), "within_10pct": sum(1 for e, a in est_pairs if e and abs((a - e) / e) <= 0.10)} if est_pairs else None},
        "collections": {"charged": round(sum(v[0] for v in coll_by_prop.values()), 2) if cd else None, "paid": round(sum(v[1] for v in coll_by_prop.values()), 2) if cd else None,
                         "rate_pct": pct(sum(v[1] for v in coll_by_prop.values()), sum(v[0] for v in coll_by_prop.values()), 2) if cd else None,
                         "by_manager": {m: {"charged": round(c, 0), "rate_pct": pct(pd_, c)} for m, (c, pd_) in sorted({mm: (sum(v[0] for p, v in coll_by_prop.items() if props.get(p, {}).get("manager") == mm), sum(v[1] for p, v in coll_by_prop.items() if props.get(p, {}).get("manager") == mm)) for mm in {props.get(p, {}).get("manager", "Unassigned") for p in coll_by_prop}}.items()) if c > 20000} if (cd and props) else None,
                         "note": f"rent charges only (GL {rent_gl}), Charge Status All; short-term stays not excluded"},
        "delinquency": {"current_tenants_with_balance": sum(d["tenants"] for d in delq_by_prop.values()) if dq_rep else None, "in_eviction": sum(d["evict"] for d in delq_by_prop.values()) if dq_rep else None,
                         "receivable_current_tenants": round(sum(d["receivable"] for d in delq_by_prop.values()), 2) if dq_rep else None, "legacy_receivable_past_tenants": round(legacy, 2) if legacy is not None else None},
        "team": {"managers": dict(collections.Counter(p["manager"] for p in props.values() if p["active"])) if props else None},
        "survey": {"responses": len(survey or []), "nps": None, "promoters": None, "passives": None, "detractors": None, "pickup_pct": None},
    }
    if survey:
        nps_vals = [s.get("nps") for s in survey if isinstance(s.get("nps"), (int, float))]
        pr = sum(1 for v in nps_vals if v >= 9); de = sum(1 for v in nps_vals if v <= 6)
        calls = [s.get("call_response") for s in survey if s.get("call_response")]
        company["survey"].update({"nps": round(100 * (pr - de) / len(nps_vals)) if nps_vals else None, "promoters": pr, "passives": len(nps_vals) - pr - de, "detractors": de,
                                  "pickup_pct": pct(sum(1 for c in calls if str(c).lower().startswith("picked")), len(calls)),
                                  "above_market_maintenance": sum(1 for s in survey if "above market" in str(s.get("maintenance_cost", "")).lower()),
                                  "statements_confusing": sum(1 for s in survey if "confusing" in str(s.get("accounting", "")).lower()),
                                  "matched_to_owner_directory": None})

    # ------------------------------------------------------------- per owner
    weights = cfg.get("risk_weights", {})
    W = lambda k, d: weights.get(k, d)
    # a lease-expiration row on a unit the rent roll shows vacant is a stale lease (tenant already gone); it must not count as an expiring lease
    vacant_names_by_prop = {p: {u["unit"] for u in us if u["status"].startswith("Vacant")} for p, us in units_by_prop.items()}
    lease_to_by_unit = {(p, u["unit"]): u["lease_to"] for p, us in units_by_prop.items() for u in us if u.get("lease_to")}
    def is_stale(p, x):   # on a vacant unit, or its expiry disagrees with the rent roll's current Lease To for that unit (the unit was re-leased; this row is the old tenant)
        if x["unit"] in vacant_names_by_prop.get(p, set()): return True
        lt = lease_to_by_unit.get((p, x["unit"])); return bool(lt and abs((lt - x["expires"]).days) > 1)
    def live_expiring(p): return [x for x in expiring_by_prop.get(p, []) if not is_stale(p, x)]
    stale_rows = sum(1 for p, xs in expiring_by_prop.items() for x in xs if is_stale(p, x))
    if stale_rows: dq.append(f"{stale_rows} lease-expiration row(s) are stale (unit vacant on the rent roll, or re-leased since); ignored for expiring-lease signals and lease-end dates.")
    OCC_STATUSES = ("Current", "Notice-Rented", "Notice-Unrented", "Evict")
    def releassed_since_safe(p, t):   # an open turn on a unit that has a tenant who moved in after the work order was opened is not really open for the owner
        return any(u["unit"] == t.get("unit") and u["status"] in OCC_STATUSES and u.get("move_in") and u["move_in"] > t["created"] for u in units_by_prop.get(p, []))
    matched = 0; placeholder_units_total = [0]
    for oid, o in owners.items():
        P = o["property_ids"]
        if not P and not o["all_property_ids"]: continue
        units = [u for p in P for u in units_by_prop.get(p, [])]
        rec = {"owner_id": oid, "first_name": o["first_name"], "name": o["name"], "display_name": pretty_name(o["name"]), "emails": o["emails"], "properties": len(P), "inactive_properties": len(o["all_property_ids"]) - len(P),
               "managers": sorted({props.get(p, {}).get("manager", "Unassigned") for p in P}), "regions": sorted({region_of(props.get(p, {}).get("city", ""), cfg) for p in P}),
               "property_types": sorted({props.get(p, {}).get("type", "") for p in P if props.get(p)}),
               "units": len(units), "occupied": sum(1 for u in units if u["status"] in OCC), "vacant": sum(1 for u in units if u["status"].startswith("Vacant")),
               "rent_month": round(sum(u["rent"] or 0 for u in units if u["status"] in OCC)), "past_due": round(sum(u["past_due"] for u in units)),
               "rent_vs_market_pct": None, "tenure_start": None, "tenure_months": None, "signals": [], "score": 0, "survey": None, "properties_detail": []}
        pairs = [(u["rent"], u["market_rent"]) for u in units if u["status"] in OCC and u["rent"] and u["market_rent"]]
        if pairs: rec["rent_vs_market_pct"] = round(100 * (sum(a for a, b in pairs) / sum(b for a, b in pairs) - 1), 1)
        starts_all = [to_dt(props[p]["management_start"]) for p in dict.fromkeys(P + o["all_property_ids"]) if props.get(p, {}).get("management_start")]
        starts_active = [to_dt(props[p]["management_start"]) for p in P if props.get(p, {}).get("management_start")]
        if starts_all: rec["tenure_start"] = min(starts_all).date().isoformat(); rec["tenure_months"] = round((TODAY - min(starts_all)).days / 30.44, 1)   # tenure counts every property ever managed, not just the active ones
        rec["active_since"] = min(starts_active).date().isoformat() if starts_active else None
        # ops signals
        def add(key, default, text, kind, points=None):
            pts = W(key, default) if points is None else points
            rec["signals"].append({"key": key, "points": pts, "signal": text, "type": kind}); rec["score"] += pts
        ending = [p for p in o["all_property_ids"] if props.get(p, {}).get("management_end") and to_dt(props[p]["management_end"]) > TODAY]
        if ending: add("management_ending", 50, f"Management end date already set on {len(ending)} property(ies)", "ops")
        exits = [m for p in dict.fromkeys(P + o["all_property_ids"]) for m in moveouts if m["property_id"] == p and m["owner_initiated"]]
        if exits: add("owner_exit_moveout", 40, f"{len(exits)} move-out(s) recorded as owner-initiated in the window", "ops")
        notes = [t for p in P for t in turns_by_prop.get(p, []) if t["owner_exit_note"]]
        if notes: add("owner_exit_note", 30, f"{len(notes)} turn(s) canceled with an owner-exit note", "ops")
        open_turns = [t for p in P for t in turns_by_prop.get(p, []) if t.get("open") and t["type"] != "Internal" and (TODAY - t["created"]).days > 45 and not releassed_since_safe(p, t)]
        if open_turns: add("turn_open_45d", 10, f"{len(open_turns)} turn work order(s) open more than 45 days (oldest {max((TODAY - t['created']).days for t in open_turns)} days)", "ops")
        loss = sum(legacy_by_prop.get(p, 0.0) for p in P)
        if loss > max(3000, rec["rent_month"] or 0): add("prior_tenant_loss", 10, f"Former tenants left ${loss:,.0f} unpaid on the owner's properties (more than a month of rent)", "ops")
        vac = [u for p in P for u in vac_by_prop.get(p, [])]
        long_vac = [u for u in vac if u["days_vacant"] and u["days_vacant"] > 60]; mid_vac = [u for u in vac if u["days_vacant"] and 30 < u["days_vacant"] <= 60]
        if long_vac: add("vacancy_60", 25, f"{len(long_vac)} unit(s) vacant more than 60 days", "ops")
        elif mid_vac: add("vacancy_30", 15, f"{len(mid_vac)} unit(s) vacant more than 30 days", "ops")
        if vac and rr is None and uv is None: pass
        dl = [delq_by_prop[p] for p in P if p in delq_by_prop]
        recv = sum(d["receivable"] for d in dl); month_rent = max(rec["rent_month"], 1)
        if recv > month_rent: add("delinquency", 10, f"Tenant balances of ${recv:,.0f} exceed a month of rent", "ops")
        exp = [x for p in P for x in live_expiring(p) if not x["renewal_started"] and not x["notice_given"] and not re.search(r"notice|move", x["status"], re.I)]
        if exp: add("expiring_no_renewal", 10, f"{len(exp)} lease(s) expiring within 120 days with no renewal started", "ops", points=min(W("expiring_no_renewal", 10) * len(exp), 20))
        if rec["tenure_months"] is not None and rec["tenure_months"] < 12: add("first_year", 10, f"First year with the company ({rec['tenure_months']:.0f} months)", "ops")
        elif rec.get("active_since") and (TODAY - to_dt(rec["active_since"])).days < 365 and rec["tenure_months"] is not None:
            add("new_property", 0, f"Long-time owner ({rec['tenure_months']:.0f} months) whose current property only started {rec['active_since']}", "ops")
        if o.get("hold_payments") and (rec["tenure_months"] or 0) >= 6: add("payments_on_hold", 15, "Owner payments on hold (a tenured owner: dispute, negative balance or owner request)", "ops")
        lp = o.get("last_payment")
        if lp is not None and rec["occupied"] and rec["rent_month"] and (rec["tenure_months"] or 0) >= 3 and (TODAY - lp).days > 90:
            add("no_owner_payment_90d", 10, f"No owner payment in {(TODAY - lp).days} days despite {rec['occupied']} leased unit(s)", "ops")
        rec["tags"] = o.get("tags", "")
        # survey signals
        resp = [s for em in o["emails"] for s in surv_by_email.get(em, [])]
        if resp:
            matched += 1; s = sorted(resp, key=lambda x: x.get("completed_at") or "")[-1]; rec["survey"] = s
            nps = s.get("nps")
            if isinstance(nps, (int, float)):
                if nps <= 6: add("detractor", 40, f"Detractor: rated {nps:.0f}/10 on likelihood to recommend", "survey")
                elif nps <= 8: add("passive", 15, f"Passive: rated {nps:.0f}/10", "survey")
                else: add("promoter", -20, f"Promoter: rated {nps:.0f}/10", "survey")
            mc = str(s.get("maintenance_cost", "")).lower()
            if "significantly above" in mc: add("cost_sig_above", 25, "Rates maintenance costs significantly above market", "survey")
            elif "above market" in mc: add("cost_above", 15, "Rates maintenance costs somewhat above market", "survey")
            if "confusing" in str(s.get("accounting", "")).lower(): add("statements_confusing", 10, "Finds the monthly statements confusing", "survey")
            if str(s.get("call_response", "")).lower().startswith("voicemail"): add("voicemail", 10, "Got voicemail the last time they called", "survey")
            if s.get("comment"): rec["signals"].append({"key": "comment", "points": 0, "signal": "Left a written comment (read it before calling)", "type": "survey"})
            if s.get("testimonial_approved") or s.get("has_testimonial"): add("testimonial", -10, "Offered a testimonial", "survey")
        # property detail for owner-level output
        for p in P:
            pr = props.get(p, {}); pu = units_by_prop.get(p, []); pv = vac_by_prop.get(p, []); pe = live_expiring(p); cc = coll_by_prop.get(p)
            ppairs = [(u["rent"], u["market_rent"]) for u in pu if u["status"] in OCC and u["rent"] and u["market_rent"]]
            vd = {u.get("unit"): u for u in pv}; ud = []
            for u in pu:
                st = u["status"]; v = vd.get(u["unit"]); kind = "vacant" if st.startswith("Vacant") else ("notice" if st.startswith("Notice") else "leased")
                never = kind == "vacant" and st == "Vacant-Unrented" and u.get("last_move_out") is None and not (v and (v.get("days_vacant") or 0) > 0)
                ud.append({"unit": u["unit"], "status": kind, "pre_leased": st == "Vacant-Rented", "never_leased": never, "rent": u["rent"], "lease_end": u["lease_to"].date().isoformat() if u.get("lease_to") else None,
                           "days_vacant": (v["days_vacant"] if v else None) if kind == "vacant" else None, "rent_ready": (v["rent_ready"] if v else None) if kind == "vacant" else None, "posted": (v["posted"] if v else None) if kind == "vacant" else None})
            # AppFolio buildings often carry floor-plan placeholder rows as units (never occupied, no move-out, 0 days vacant). On a building with tenants they are
            # not real vacancies and must not inflate the owner's unit count; on a building with nothing leased yet (lease-up) every unit is legitimately unleased, so keep them.
            leased_any = any(x["status"] != "vacant" for x in ud)
            placeholders = [x for x in ud if x["never_leased"]] if leased_any else []
            if placeholders: ud = [x for x in ud if not x["never_leased"]]; placeholder_units_total[0] += len(placeholders)
            rec["properties_detail"].append({"property_id": p, "type": pr.get("type", ""), "city": pr.get("city", ""), "address": pr.get("address", ""), "units_detail": ud, "placeholder_units": len(placeholders), "zip": pr.get("zip", ""), "manager": pr.get("manager", "Unassigned"), "managed_since": pr.get("management_start"),
                                             "units": len(ud), "occupied": sum(1 for u in pu if u["status"] in OCC), "rent_month": round(sum(u["rent"] or 0 for u in pu if u["status"] in OCC)),
                                             "rent_vs_market_pct": round(100 * (sum(a for a, b in ppairs) / sum(b for a, b in ppairs) - 1), 1) if ppairs else None, "past_due": round(sum(u["past_due"] for u in pu)),
                                             "next_lease_end": min((x["expires"].date().isoformat() for x in pe), default=None), "lease_ends_120d": sorted(x["expires"].date().isoformat() for x in pe), "expiring_120d": len(pe),
                                             "m2m_units": sum(1 for un in m2m_units_by_prop.get(p, []) if un not in vacant_names_by_prop.get(p, set())),   # a month-to-month row on a vacant unit is a stale lease
                                             # vacancy figures come from the cleaned unit list (placeholders dropped), so the tile and the sentence always agree
                                             "vacant_units": sum(1 for x in ud if x["status"] == "vacant"), "max_days_vacant": max((x["days_vacant"] or 0 for x in ud if x["status"] == "vacant"), default=None) if any(x["status"] == "vacant" for x in ud) else None,
                                             "rent_ready": any(x.get("rent_ready") for x in ud if x["status"] == "vacant") if any(x["status"] == "vacant" for x in ud) else None, "lease_up": bool(ud) and all(x["status"] == "vacant" and x["never_leased"] for x in ud),
                                             "turns_12mo": len(turns_by_prop.get(p, [])), "turn_median_days": med([t["days"] for t in turns_by_prop.get(p, [])]),
                                             "collected_90d_pct": pct(cc[1], cc[0], 1) if cc and cc[0] else None, "charged_90d": round(cc[0]) if cc else None})
        thr = cfg.get("risk_tiers", {"call": 50, "watch": 25})
        if not P and o["all_property_ids"]:
            rec["tier"] = "Departed"  # every property has a past management end date: already gone, count it, do not call it
        else:
            rec["tier"] = "Call this week" if rec["score"] >= thr["call"] else ("Watch" if rec["score"] >= thr["watch"] else ("Healthy" if rec["signals"] else "No signal"))
        rec["signals"].sort(key=lambda x: -abs(x["points"]))
        rec["units"] = sum(pd["units"] for pd in rec["properties_detail"]); rec["vacant"] = rec["units"] - rec["occupied"]
        rec["departed_recent"] = rec["tier"] == "Departed" and any(props.get(p, {}).get("management_end") and (TODAY - to_dt(props[p]["management_end"])).days <= 365 for p in o["all_property_ids"])
        out_owners[oid] = rec
    if placeholder_units_total[0]: dq.append(f"{placeholder_units_total[0]} rent-roll unit rows on tenanted buildings look like floor-plan placeholders (never occupied, no move-out, 0 days vacant); left out of owner-facing unit counts.")
    company["survey"]["matched_to_owner_directory"] = matched if survey else None
    if survey and owners and matched < len(survey) * 0.5: dq.append(f"Only {matched} owners matched a survey response by email; check that the owner directory emails match the survey roster.")
    thr = cfg.get("risk_tiers", {"call": 50, "watch": 25})
    def tier_for(score, sigs, all_departed):
        if all_departed: return "Departed"
        return "Call this week" if score >= thr["call"] else ("Watch" if score >= thr["watch"] else ("Healthy" if sigs else "No signal"))
    def merge_by_inbox(recs):
        """One record per inbox. AppFolio often holds one person as several owner records (a person plus LLCs, or two records on one property);
        scoring each separately double-counts a shared building and hands the manager two cards for one call. Signals de-duplicate by key, score and tier are recomputed."""
        groups, singles, out = {}, [], []
        for r in recs:
            em = r["emails"][0].lower() if r["emails"] else None
            (groups.setdefault(em, []) if em else singles).append(r)
        for em, grp in groups.items():
            if len(grp) == 1: out.append(grp[0]); continue
            P, seen, best = [], set(), {}
            for g in grp:
                for pd in g["properties_detail"]:
                    if pd["property_id"] not in seen: seen.add(pd["property_id"]); P.append(pd)
                for sg in g["signals"]:
                    k = sg.get("key") or sg["signal"]
                    if k not in best or abs(sg["points"]) > abs(best[k]["points"]): best[k] = sg
            tenure_all = max((g["tenure_months"] for g in grp if g.get("tenure_months") is not None), default=None)
            if tenure_all is not None and tenure_all >= 12: best.pop("first_year", None)   # the merged person is not new; one of the records only looked new
            sigs = sorted(best.values(), key=lambda x: -abs(x["points"])); score = sum(x["points"] for x in sigs)
            persons = [g for g in grp if g["first_name"] and not ENTITY_RE.search(g["name"]) and not re.match(r"^[#\d]", g["name"].strip()) and g["name"].count(",") == 1]
            first = (persons or [g for g in grp if g["first_name"]] or [grp[0]])[0]["first_name"]
            units = [u for pd in P for u in pd["units_detail"]]
            m = dict(grp[0]); m.update({
                "owner_id": "+".join(g["owner_id"] for g in grp), "name": " / ".join(dict.fromkeys(g["name"] for g in grp)), "first_name": first, "merged_from": [pretty_name(g["name"]) for g in grp],
                "display_name": pretty_name((persons or sorted([g for g in grp if not re.match(r"^[#\d]", g["name"].strip())] or grp, key=lambda g: len(g["name"])))[0]["name"]),
                "emails": list(dict.fromkeys(e for g in grp for e in g["emails"])), "properties_detail": P, "properties": len(P), "inactive_properties": sum(g["inactive_properties"] for g in grp),
                "managers": list(dict.fromkeys(mm for g in grp for mm in g["managers"])), "regions": sorted({x for g in grp for x in g["regions"]}), "property_types": sorted({x for g in grp for x in g["property_types"]}),
                "units": len(units), "occupied": sum(1 for u in units if u["status"] != "vacant"), "vacant": sum(1 for u in units if u["status"] == "vacant"),
                "rent_month": sum(pd.get("rent_month") or 0 for pd in P), "past_due": sum(pd.get("past_due") or 0 for pd in P),
                "tenure_start": min((g["tenure_start"] for g in grp if g.get("tenure_start")), default=None), "tenure_months": max((g["tenure_months"] for g in grp if g.get("tenure_months") is not None), default=None),
                "signals": sigs, "score": score, "tier": tier_for(score, sigs, all(g["tier"] == "Departed" for g in grp)),
                "survey": min((g["survey"] for g in grp if g.get("survey")), key=lambda sv: sv.get("nps", 10) if isinstance(sv.get("nps"), (int, float)) else 10, default=None),
                "tags": "; ".join(t for t in dict.fromkeys(g.get("tags", "") for g in grp) if t), "departed_recent": any(g.get("departed_recent") for g in grp),
                "active_since": min((g["active_since"] for g in grp if g.get("active_since")), default=None)})
            out.append(m)
        return out + singles
    ranked = merge_by_inbox(list(out_owners.values()))
    by_prop = collections.defaultdict(list)
    for r in ranked:
        for pd in r["properties_detail"]: by_prop[pd["property_id"]].append(r)
    for r in ranked:
        r["co_owners"] = [{"name": other["name"], "address": pd.get("address") or pd.get("city") or pd["property_id"]} for pd in r["properties_detail"] for other in by_prop[pd["property_id"]] if other is not r]
    company["departed_owners"] = {"count": sum(1 for r in ranked if r["tier"] == "Departed"), "recent_12mo": sum(1 for r in ranked if r.get("departed_recent"))}   # after the inbox merge, so it matches the tiles
    merged_n = sum(1 for r in ranked if r.get("merged_from"))
    if merged_n: dq.append(f"{merged_n} inbox(es) held more than one owner record; merged into one record each (see merged_from).")
    company["watch_tags"] = [{"label": t.get("label", t["match"]), "match": t["match"], "owners": sum(1 for r in ranked if r["tier"] != "Departed" and t["match"].lower() in (r.get("tags") or "").lower()),
                              "at_risk": sum(1 for r in ranked if r["tier"] in ("Call this week", "Watch") and t["match"].lower() in (r.get("tags") or "").lower())} for t in cfg.get("watch_tags", []) if t.get("match")]
    ranked.sort(key=lambda r: -r["score"])
    return company, ranked, dq

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    a = sub.add_parser("inventory"); a.add_argument("exports_dir")
    a.add_argument("--json", action="store_true", help="machine-readable output")
    b = sub.add_parser("features"); b.add_argument("exports_dir"); b.add_argument("--out", required=True); b.add_argument("--survey"); b.add_argument("--as-of"); b.add_argument("--config"); b.add_argument("--json", action="store_true", help="machine-readable output")
    args = ap.parse_args()
    reports = load_dir(args.exports_dir)
    inv = {k: {"file": os.path.basename(v.path), "rows": len(v.rows), "exported_on": v.as_of, "columns": len(v.hdr)} for k, v in reports.items()}
    missing = [k for k in SIGNATURES if k not in reports]
    if args.cmd == "inventory":
        if args.json: print(json.dumps({"found": inv, "missing": missing}, indent=1)); return
        print("Recognised: " + ", ".join(f"{k} ({v['exported_on']})" for k, v in sorted(inv.items())) + ".")
        print("Missing: " + (", ".join(missing) if missing else "none") + ("" if not missing else ". See references/appfolio-exports.md for the export steps.")); return
    cfg = json.load(open(args.config)) if args.config else {}
    survey = json.load(open(args.survey)) if args.survey else None
    if isinstance(survey, dict): survey = survey.get("responses", [])
    as_of = args.as_of or max((v.as_of for v in reports.values()), default=dt.date.today().isoformat())
    company, owners, dq = build(reports, cfg, survey, as_of)
    os.makedirs(args.out, exist_ok=True)
    json.dump({"found": inv, "missing": missing, "as_of": as_of}, open(os.path.join(args.out, "inventory.json"), "w"), indent=1)
    json.dump(company, open(os.path.join(args.out, "company.json"), "w"), indent=1, default=str)
    json.dump(owners, open(os.path.join(args.out, "owners.json"), "w"), indent=1, default=str)
    with open(os.path.join(args.out, "data-quality.md"), "w") as f:
        f.write(f"# Data quality notes (as of {as_of})\n\n")
        f.write("Reports found: " + ", ".join(sorted(reports)) + "\n\nMissing: " + (", ".join(missing) or "none") + "\n\n")
        for n in dq: f.write(f"- {n}\n")
    tiers = collections.Counter(o["tier"] for o in owners)
    summary = {"as_of": as_of, "reports": sorted(reports), "missing": missing, "owners": len(owners), "tiers": dict(tiers), "survey_responses": len(survey or []), "notes": dq}
    json.dump(summary, open(os.path.join(args.out, "features-summary.json"), "w"), indent=1)
    if args.json: print(json.dumps(summary, indent=1))
    else: print(f"Features built from {len(reports)} reports as of {as_of} (missing: {', '.join(missing) or 'none'}); {len(owners)} owner records; {len(survey or [])} survey responses; {len(dq)} data-quality notes in data-quality.md.")

if __name__ == "__main__":
    main()
